"""23E 模板模式引擎（2026-10-06 装配工作流·能力层）。

以《MBA学位论文模板格式（23E自制）.docx》为**样式库**装配整篇论文：
- 模板只作样式库：剥离其正文内容与节残留（页眉"参考文献"、pgNumType 等），
  节结构由引擎按单元重建；
- 每单元一个节（NEW_PAGE）：摘要/Abstract/目录用罗马页码（自摘要起 I），
  第1章起阿拉伯页码重新计 1；页眉＝单元标题；
- 目录为 Word TOC 域（\\o "1-3" \\h \\z \\u），章/一级节/二级节样式写入
  outlineLvl（同时惠及导航窗格）；摘要/Abstract/目录标题用克隆样式
  「MBA论文前置标题」（无 outline，不进目录）；
- 章、节、图表题一律**照排原文**（md 已带"第N章"/"图X.Y"/"表X.Y"编号，
  不再自动编号）。

仅新增文件与新增路径，FormatEngine 旧模式行为不变。
"""

from __future__ import annotations

import copy
import re
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from format_engine import (
    FormatEngine, _set_font, _set_paragraph_spacing, _add_field_code,
    _set_header_footer,
    FONT_BODY_CN, FONT_BODY_EN, FONT_ABSTRACT_TITLE_EN,
    SIZE_BODY, SIZE_CAPTION, SIZE_CHAPTER,
    LINE_SPACING_BODY, LINE_SPACING_SINGLE, TWO_CHAR_INDENT,
    SPACING_CHAPTER_BEFORE, SPACING_CHAPTER_AFTER,
    A4_WIDTH, A4_HEIGHT,
)

# 23E 模板样式名映射（模板须含下列样式，缺失即报错）
STYLE = {    "chapter": "MBA论文章标题",
    "h1": "MBA论文一级节标题",
    "h2": "MBA论文二级节标题",
    "h3": "MBA论文三级节标题",
    "body": "MBA论文正文",
    "fig_caption": "MBA论文图标题",
    "tab_caption": "MBA表标题",
    "note": "MBA论文图表文字",
    "reference": "MBA论文参考文献",
    "front_title": "MBA论文前置标题",   # 由章标题克隆（无 outline），摘要/Abstract/目录用
    "appendix_h1": "MBA论文附录节标题",  # 由一级节标题克隆（无 outline），附录 A.1/B.1 等用（2026-10-06 裁定：附录小节不进大纲/目录）
    "table": "MBA论文表格",             # 表样式（可选，存在则套用）
}

# outline 级别：章→0、一级节→1、二级节→2（TOC \o "1-3" 与导航窗格）
_OUTLINE = {"chapter": 0, "h1": 1, "h2": 2}

_FRONT_UNITS = {"abstract_cn": "摘　要", "abstract_en": "Abstract", "toc": "目　录"}

# 编号与标题之间的分隔：半角空格 → 全角空格（2026-10-06 用户裁定；
# md 源文件保持半角不动，转换只发生在渲染层与 docx 修复层）
_SEP_RES = [
    re.compile(r"^(第\d+章)\s+"),                                # 章号
    re.compile(r"^([A-Z0-9]+(?:\.[A-Z0-9]+)+)\s+"),              # 节号（1.1／3.3.2／A.1）
    re.compile(r"^((?:附?表|图)[A-Z0-9]+(?:\.[A-Z0-9]+)?)\s+"),  # 表号/图号（表3.2／附表A.1／图1.1）
]


def _fullwidth_sep(text: str) -> str:
    """章号/节号/表号/图号与标题之间的半角空格串 → 一个全角空格。"""
    for pat in _SEP_RES:
        m = pat.match(text)
        if m:
            return m.group(1) + "　" + text[m.end():]
    return text


class Thesis23EEngine(FormatEngine):
    """以 23E 模板为样式库的论文装配引擎。"""

    def __init__(self, template_path: str, landscape_tables=()):
        # 不调 super().__init__：不建空文档、不定义 NJU 样式
        self.mode = "thesis"
        self.doc = Document(template_path)
        self._chapter_num = 0
        self._figure_num = 0
        self._table_num = 0
        self._current_h1_text = ""
        self._first_unit_done = False
        self._appendix = False            # 附录单元渲染时置 True（小节标题去大纲）
        # 横排宽表清单（用户手工横排沉淀；键为题注前缀，如 "表3.3"）：
        # 渲染到该表题注时自动开横排节，题注+表+来源注在节内，注毕收回纵排
        self._landscape_tables = set(landscape_tables)
        self._landscape_open = False
        self._strip_template_body()
        self._prepare_styles()

    # ── 横排节（宽表）─────────────────────────────────────────
    def _is_landscape_caption(self, text: str) -> bool:
        norm = text.replace("　", "").replace(" ", "")
        return any(norm.startswith(p) for p in self._landscape_tables)

    def _landscape_on(self):
        from docx.enum.section import WD_ORIENT
        sec = self.doc.add_section(WD_SECTION_START.NEW_PAGE)
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = A4_HEIGHT, A4_WIDTH
        self._set_pgnum(sec)   # 清掉克隆来的 pgNumType（防页码重计）
        self._landscape_open = True

    def _landscape_off(self):
        from docx.enum.section import WD_ORIENT
        sec = self.doc.add_section(WD_SECTION_START.NEW_PAGE)
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = A4_WIDTH, A4_HEIGHT
        self._set_pgnum(sec)
        self._landscape_open = False

    # ── 模板预处理 ───────────────────────────────────────────
    def _strip_template_body(self):
        """删除模板全部正文内容，仅保留 body 级 sectPr（页面设置），
        并清掉节残留（页眉页脚引用、页码格式）。"""
        body = self.doc.element.body
        for child in list(body):
            if child.tag == qn("w:sectPr"):
                for junk in child.findall(qn("w:headerReference")) + \
                            child.findall(qn("w:footerReference")) + \
                            child.findall(qn("w:pgNumType")):
                    child.remove(junk)
            else:
                body.remove(child)

    # CT_Style 中排在 uiPriority 之后的元素（uiPriority 必须位于它们之前；
    # WPS 存盘的模板常把 uiPriority 写错位置，Word 严格校验会报"无法读取的内容"）
    _STYLE_AFTER_UIPRIORITY = ("w:semiHidden", "w:unhideWhenUsed", "w:qFormat",
                               "w:locked", "w:personal", "w:personalCompose",
                               "w:personalReply", "w:rsid", "w:pPr", "w:rPr",
                               "w:tblStylePr", "w:tblPr", "w:tcPr", "w:trPr")

    def _prepare_styles(self):
        """校验必需样式、写 outlineLvl、克隆前置标题样式、归一化样式元素顺序。"""
        missing = [v for k, v in STYLE.items()
                   if k not in ("front_title", "appendix_h1", "table")
                   and v not in self.doc.styles]
        if missing:
            raise KeyError(f"23E 模板缺少样式：{missing}")
        # 模板（WPS 存盘）样式顺序归一化：uiPriority 移到 schema 位置
        for st in self.doc.styles.element.findall(qn("w:style")):
            uip = st.find(qn("w:uiPriority"))
            if uip is None:
                continue
            for tag in self._STYLE_AFTER_UIPRIORITY:
                nxt = st.find(qn(tag))
                if nxt is not None:
                    # 已在正确位置（uiPriority 于其后元素之前）则跳过
                    if list(st).index(uip) > list(st).index(nxt):
                        st.remove(uip)
                        nxt.addprevious(uip)
                    break
        for key, lvl in _OUTLINE.items():
            style = self.doc.styles[STYLE[key]]
            pPr = style.element.get_or_add_pPr()
            ol = pPr.find(qn("w:outlineLvl"))
            if ol is None:
                ol = OxmlElement("w:outlineLvl")
                pPr.append(ol)
            ol.set(qn("w:val"), str(lvl))
        # 克隆章标题 → 前置标题（去 outline，摘要/Abstract/目录不进 TOC）
        src = self.doc.styles[STYLE["chapter"]]
        new_el = copy.deepcopy(src.element)
        new_el.set(qn("w:styleId"), "MBAFrontTitle")
        new_el.find(qn("w:name")).set(qn("w:val"), STYLE["front_title"])
        pPr = new_el.find(qn("w:pPr"))
        if pPr is not None:
            ol = pPr.find(qn("w:outlineLvl"))
            if ol is not None:
                pPr.remove(ol)
        self.doc.styles.element.append(new_el)
        # 克隆一级节标题 → 附录节标题（去 outline，附录 A.1/B.1/C.1 等不进 TOC）
        src = self.doc.styles[STYLE["h1"]]
        new_el = copy.deepcopy(src.element)
        new_el.set(qn("w:styleId"), "MBAAppendixH1")
        new_el.find(qn("w:name")).set(qn("w:val"), STYLE["appendix_h1"])
        pPr = new_el.find(qn("w:pPr"))
        if pPr is not None:
            ol = pPr.find(qn("w:outlineLvl"))
            if ol is not None:
                pPr.remove(ol)
        self.doc.styles.element.append(new_el)

    # ── 节管理 ───────────────────────────────────────────────
    # ECMA-376 CT_SectPr 顺序中排在 w:pgNumType 之后的元素（pgNumType 必须插到它们之前，
    # 否则 Word 报"无法读取的内容"——2026-10-06 因此返修一次）
    _SECTPR_AFTER_PGNUM = ("w:cols", "w:formProt", "w:vAlign", "w:noEndnote",
                           "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter",
                           "w:docGrid", "w:printerSettings", "w:sectPrChange")

    @staticmethod
    def _set_pgnum(section, fmt: str | None = None, start: int | None = None):
        """设置节的页码格式；总是先清掉从上一节克隆来的 pgNumType。"""
        sectPr = section._sectPr
        for old in sectPr.findall(qn("w:pgNumType")):
            sectPr.remove(old)
        if fmt is None and start is None:
            return
        pg = OxmlElement("w:pgNumType")
        if fmt:
            pg.set(qn("w:fmt"), fmt)
        if start is not None:
            pg.set(qn("w:start"), str(start))
        for tag in Thesis23EEngine._SECTPR_AFTER_PGNUM:
            nxt = sectPr.find(qn(tag))
            if nxt is not None:
                nxt.addprevious(pg)
                return
        sectPr.append(pg)

    def _start_section(self, header_text: str, fmt: str, start: int | None = None):
        """开新节（NEW_PAGE）＋页眉单元标题＋页脚页码。内容须在此调用之后添加。"""
        sec = self.doc.add_section(WD_SECTION_START.NEW_PAGE)
        _set_header_footer(sec, header_text)
        self._set_pgnum(sec, fmt, start)
        return sec

    # ── 标题与段落写入 ────────────────────────────────────────
    def _add_unit_title(self, text: str, front: bool = False):
        text = _fullwidth_sep(text)
        style = STYLE["front_title"] if front else STYLE["chapter"]
        para = self.doc.add_paragraph(text, style=style)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_paragraph_spacing(para, SPACING_CHAPTER_BEFORE, SPACING_CHAPTER_AFTER,
                               LINE_SPACING_SINGLE, WD_ALIGN_PARAGRAPH.CENTER)
        return para

    def add_body(self, text: str):  # noqa: D401 - 覆盖：用模板样式，首行缩进由样式承担
        para = self.doc.add_paragraph(style=STYLE["body"])
        for r in para.runs:
            r._element.getparent().remove(r._element)
        self._parse_inline(text, para)

    def add_image(self, image_path: str, width_inches: float = 5.5):
        from docx.shared import Inches as DocInches
        para = self.doc.add_paragraph(style=STYLE["body"])
        para.paragraph_format.first_line_indent = Pt(0)
        _set_paragraph_spacing(para, Pt(6), Pt(6), LINE_SPACING_SINGLE, WD_ALIGN_PARAGRAPH.CENTER)
        run = para.add_run()
        run.add_picture(image_path, width=DocInches(width_inches))

    def add_display_math(self, latex: str, eq_num: str | None = None):
        """独立行 LaTeX \\[...\\] → OMML（经 pandoc），居中；模板无 NJU Body 样式，改用 23E 正文样式。"""
        import copy
        ommls = self._latex_to_omml(latex, display=True)
        if not ommls:
            self.add_body(f"[Equation: {latex}]")
            return
        para = self.doc.add_paragraph(style=STYLE["body"])
        para.paragraph_format.first_line_indent = Pt(0)
        _set_paragraph_spacing(para, Pt(6), Pt(6), LINE_SPACING_SINGLE, WD_ALIGN_PARAGRAPH.CENTER)
        for omml in ommls:
            para._element.append(copy.deepcopy(omml))

    def _add_caption(self, el, style_name: str):
        """图表题照排：label 粗体居中，rest（附表说明）非粗体接续。"""
        para = self.doc.add_paragraph(style=style_name)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent = Pt(0)
        run = para.add_run(_fullwidth_sep(el["text"]))
        _set_font(run, FONT_BODY_CN, FONT_BODY_EN, SIZE_CAPTION, bold=True)
        if el.get("rest"):
            run2 = para.add_run(el["rest"])
            _set_font(run2, FONT_BODY_CN, FONT_BODY_EN, SIZE_CAPTION, bold=False)
        return para

    def _add_source_note(self, text: str):
        """图注/表注：全部版式由样式承担（规范 3.6(6)(7)：左缩进两字符＋悬挂
        两字符、两端对齐、10.5pt、单倍行距、段前后 0），不写段落级覆盖。"""
        para = self.doc.add_paragraph(style=STYLE["note"])
        for r in para.runs:
            r._element.getparent().remove(r._element)
        self._parse_inline(text, para, FONT_BODY_CN, FONT_BODY_EN, SIZE_CAPTION)
        return para

    def _add_list_item(self, text: str):
        para = self.doc.add_paragraph(style=STYLE["body"])
        para.paragraph_format.first_line_indent = Pt(0)
        run = para.add_run("• ")
        _set_font(run, FONT_BODY_CN, FONT_BODY_EN, SIZE_BODY)
        self._parse_inline(text, para)
        return para

    def _add_keywords(self, keywords: list[str], en: bool = False):
        para = self.doc.add_paragraph(style=STYLE["body"])
        if en:
            label, sep = "Key Words: ", ", "
        else:
            label, sep = "关键词：", "；"
        run = para.add_run(label)
        _set_font(run, FONT_BODY_CN if not en else FONT_ABSTRACT_TITLE_EN,
                  FONT_BODY_EN, SIZE_BODY, bold=True)
        run2 = para.add_run(sep.join(keywords))
        _set_font(run2, FONT_BODY_CN, FONT_BODY_EN, SIZE_BODY)
        return para

    def _add_reference_entry(self, text: str):
        """参考文献条目照排（编号 [N] 由文献表自带，不重排）。"""
        para = self.doc.add_paragraph(style=STYLE["reference"])
        run = para.add_run(text)
        _set_font(run, FONT_BODY_CN, FONT_BODY_EN, SIZE_CAPTION)
        return para

    def add_three_line_table_23e(self, headers: list[str], rows: list[list[str]]):
        table = self.add_three_line_table(headers, rows)  # 复用旧实现（边框/字体已内置）
        try:
            table.style = self.doc.styles[STYLE["table"]]
        except KeyError:
            pass
        return table

    # ── 单元渲染 ─────────────────────────────────────────────
    def _render_elements(self, elements: list[dict]):
        for el in elements:
            t = el["type"]
            # 横排宽表状态机：题注入横排节 → 表（留）→ 来源注毕收节
            if t == "caption_tab" and self._is_landscape_caption(el["text"]) \
                    and not self._landscape_open:
                self._landscape_on()
            elif self._landscape_open and t not in ("table", "source_note"):
                self._landscape_off()   # 表后无注的情况兜底收节
            if t == "chapter":
                # 章/附录标题已在单元级渲染，正文中不应再出现；防御性按一级节处理
                self._render_h1(el["text"])
            elif t == "h1":
                self._render_h1(el["text"])
            elif t == "h2":
                para = self.doc.add_paragraph(_fullwidth_sep(el["text"]), style=STYLE["h2"])
            elif t == "h3":
                para = self.doc.add_paragraph(_fullwidth_sep(el["text"]), style=STYLE["h3"])
            elif t == "body":
                self.add_body(el["text"])
            elif t == "image":
                self.add_image(el["src"])
            elif t == "table":
                self.add_three_line_table_23e(el["headers"], el["rows"])
            elif t == "caption_fig":
                self._add_caption(el, STYLE["fig_caption"])
            elif t == "caption_tab":
                self._add_caption(el, STYLE["tab_caption"])
            elif t == "source_note":
                self._add_source_note(el["text"])
                if self._landscape_open:
                    self._landscape_off()
            elif t == "list_item":
                self._add_list_item(el["text"])
            elif t in ("keywords_cn", "keywords_en"):
                self._add_keywords(el["keywords"], en=(t == "keywords_en"))
            elif t == "blockquote":
                self.add_body(el["text"])
            elif t == "display_math":
                self.add_display_math(el["text"], el.get("eq_num"))
            # footnote 由 footnote_inserter.py 后处理，不在此渲染
        if self._landscape_open:
            self._landscape_off()       # 单元末尾兜底收节

    def _render_h1(self, text: str):
        # 附录单元的小节标题（A.1/B.1/C.1 等）用去大纲样式，不进目录（2026-10-06 裁定）
        style = STYLE["appendix_h1"] if getattr(self, "_appendix", False) else STYLE["h1"]
        self.doc.add_paragraph(_fullwidth_sep(text), style=style)

    def render_unit(self, unit: dict):
        """渲染一个合并稿单元（节标题＝单元标题，页码规则分前置/正文）。"""
        kind, uid = unit["kind"], unit["id"]
        elements = unit["elements"]

        if kind == "abstract_cn":
            if not self._first_unit_done:
                # 摘要落 section 0（封面节由 merge_cover 前置，页码不受影响）
                _set_header_footer(self.doc.sections[0], "摘要")
                self._set_pgnum(self.doc.sections[0], "upperRoman", start=1)
                self._first_unit_done = True
            self._add_unit_title(_FRONT_UNITS["abstract_cn"], front=True)
            self._render_elements(elements)
        elif kind == "abstract_en":
            self._start_section("Abstract", "upperRoman")
            para = self._add_unit_title(_FRONT_UNITS["abstract_en"], front=True)
            for run in para.runs:  # 英文摘要标题 Arial
                run.font.name = FONT_ABSTRACT_TITLE_EN
            self._render_elements(elements)
        elif kind == "references":
            self._start_section("参考文献", "decimal")
            self._add_unit_title("参考文献")
            for el in elements:
                if el["type"] == "body":
                    self._add_reference_entry(el["text"])
        elif kind == "acknowledgement":
            self._start_section("致　谢", "decimal")
            self._add_unit_title("致　谢")
            self._render_elements(elements)
        else:  # chapter / appendix：标题取 md 原文（"第N章 …"／"附录X：…"）
            title = ""
            if elements and elements[0]["type"] == "chapter":
                title = elements[0]["text"]
                elements = elements[1:]
            start = 1 if title.startswith("第1章") else None
            self._start_section(_fullwidth_sep(title), "decimal", start=start)
            self._add_unit_title(title)
            self._appendix = (kind == "appendix")
            self._render_elements(elements)
            self._appendix = False

    def add_toc_unit(self):
        """目录单元：Word TOC 域（打开后 F9 更新；页码差异不计入一致性）。"""
        self._start_section("目　录", "upperRoman")
        self._add_unit_title(_FRONT_UNITS["toc"], front=True)
        para = self.doc.add_paragraph()
        para.paragraph_format.first_line_indent = Pt(0)
        _add_field_code(para, ' TOC \\o "1-3" \\h \\z \\u ')
