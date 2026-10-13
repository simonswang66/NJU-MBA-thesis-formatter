#!/usr/bin/env python3
"""章级增量更新器（2026-10-06 装配工作流·能力层；2026-10-07 段落级差分）。

默认走**段落级最小变更集**：对 md 元素序列与 docx 段落序做 SequenceMatcher 差分，
只重建变化/新增/删除的段落；未变段落（含 run 级格式与用户手工精调）原样保留；
分节边界与页眉引用只在结构真变化（横排表集合变化）时才动——此时自动回退
整单元重建，亦可用 --full-unit 强制。

整单元重建路径（--full-unit／自动回退）原理：
- 每个单元＝一个节：单元标题段（MBA论文章标题／MBA论文前置标题）起，
  到本节的收尾 sectPr 段止（sectPr 段保留——它承载页眉与页码设置）；
- 删除区间内全部段落与表格，用 Thesis23EEngine 的同一渲染器在该节内
  重渲染单元内容（图片关系直接在目标文档内创建，无跨文档引用问题）；
- 单元标题文字若变化，同步更新该节页眉；
- 目录是 Word 域，不在合并稿单元内，永不触碰；页码/排版差异不计入一致性。

约定（2026-10-06 用户裁定；2026-10-07 增补增量路径）：
- docx 变更前**必须**时间戳备份入 `<docx 同目录>/bak/`（--no-backup 可关）；
  全部单元无差异时不写盘、不备份；
- 表格：冻结表整张不动（md 侧有差异只告警）；非冻结表保留骨架、只写差异文字
  （结构变化即增删行列才整表重渲染并告警）；
- 附录问卷等在 docx 内被手工拆行的段落：变更区间两侧**拼接后内容一致**即
  整体保留，不按 md 段落边界重排。

用法：
    python3 section_replacer.py <论文.docx> <合并稿.md> --unit ch3 [--unit ch4]
    python3 section_replacer.py <论文.docx> <合并稿.md> --unit ch3 --dry-run
    python3 section_replacer.py <论文.docx> <合并稿.md> --unit ch3 --full-unit
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import re
import shutil
import sys
import unicodedata
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from content_parser import parse_merged_thesis
from thesis23e_engine import Thesis23EEngine, STYLE, _fullwidth_sep, _fill_title_text

# 前置单元在 docx 中的标题（引擎生成，与合并稿单元 id 的映射）
_FRONT_TITLE = {
    "abstract_cn": "摘　要",
    "abstract_en": "Abstract",
    "references": "参考文献",
    "acknowledgement": "致　谢",
}

# 表题前缀（表3.2／附表A.1）；md 与 docx 两侧共用
_TABLE_PREFIX_RE = re.compile(r"^(附?表(?:[A-Z]\d*(?:\.\d+)*|\d+(?:\.\d+)*))")

_VML_IMAGEDATA = "{urn:schemas-microsoft-com:vml}imagedata"


def norm(s: str) -> str:
    """内容级归一化（两侧同施）：NFKC、去 md 粗体、<br>→空格、空白折叠、
    非 ASCII 字符邻接空格抹除、去行首列表符（「• 」与参考文献「[n] 」字面序号）。
    check_docx_text 从本模块导入本函数——核查与差分同一口径，改则两处同步失效。"""
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("**", "")
    s = re.sub(r"<br\s*/?>", " ", s, flags=re.IGNORECASE)
    s = re.sub(r"\s*/\s*", "/", s)          # 斜杠两侧空格（用户格内换行伪差异）
    s = re.sub(r"\s+", " ", s)
    s = re.sub(r"\s+(?=[^\x00-\x7f])", "", s)
    s = re.sub(r"(?<=[^\x00-\x7f])\s+", "", s)
    s = re.sub(r"^•\s*", "", s)
    s = re.sub(r"^\[\d+\]\s*", "", s)       # 参考文献行首字面序号（序号由样式自动编号承担）
    return s.strip()


def _unit_title(unit: dict) -> str:
    """合并稿单元 → docx 中的单元标题文字。"""
    if unit["id"] in _FRONT_TITLE:
        return _FRONT_TITLE[unit["id"]]
    for el in unit["elements"]:
        if el["type"] == "chapter":
            return el["text"]
    raise ValueError(f"单元 {unit['id']} 找不到标题")


class SectionReplacer:
    def __init__(self, docx_path: str, frozen_tables: list[str] | None = None,
                 landscape_tables: list[str] | None = None):
        self.path = docx_path
        self.doc = Document(docx_path)
        # 冻结表（用户手工定版、"以后不要动"的表，如 表5.2/表6.1/表6.2）：
        # 单元替换时按题注锚点保留旧表整体（含格式与回车换行），只更新周边文字
        self.frozen_tables = frozen_tables or []
        # 渲染器复用 Thesis23EEngine 的方法，但作用于现有文档（不做模板预处理）
        self.r = Thesis23EEngine.__new__(Thesis23EEngine)
        self.r.doc = self.doc
        self.r._landscape_tables = set(landscape_tables or [])
        self.r._landscape_open = False

    # ── 区域定位 ─────────────────────────────────────────────
    _TITLE_STYLES = {"MBA论文章标题", "MBA论文前置标题"}

    @staticmethod
    def _norm_title(s: str) -> str:
        """标题匹配归一：空格/全角空格/断行一律抹除（渲染层会把编号后空格转全角、
        把 <br> 渲染为段内换行——docx 侧取文得 \\n——比较时不可直接用原文）。"""
        return (s.replace("<br>", "").replace("\n", "").replace("\r", "")
                .replace(" ", "").replace("　", "").strip())

    # ── 页眉/页脚引用暂存与重挂 ───────────────────────────────
    _REF_TAGS = ("w:headerReference", "w:footerReference")

    @classmethod
    def _sectPr_of(cls, el):
        """段（或 sectPr 本身）→ 其 sectPr 元素；无则 None。"""
        if el is None:
            return None
        if el.tag == qn("w:sectPr"):
            return el
        if el.tag != qn("w:p"):
            return None
        pPr = el.find(qn("w:pPr"))
        return pPr.find(qn("w:sectPr")) if pPr is not None else None

    @classmethod
    def _clone_refs(cls, el) -> list:
        """元素 sectPr 内的页眉/页脚引用深拷贝（区域删除前暂存用）。"""
        sp = cls._sectPr_of(el)
        if sp is None:
            return []
        return [copy.deepcopy(r) for tag in cls._REF_TAGS for r in sp.findall(qn(tag))]

    @classmethod
    def _strip_refs(cls, el):
        sp = cls._sectPr_of(el)
        if sp is None:
            return
        for tag in cls._REF_TAGS:
            for r in sp.findall(qn(tag)):
                sp.remove(r)

    @classmethod
    def _attach_refs(cls, el, refs: list):
        """引用挂回元素 sectPr（先剥后挂）；headerReference/footerReference 位于 CT_SectPr 序首。"""
        sp = cls._sectPr_of(el)
        if sp is None or not refs:
            return
        cls._strip_refs(el)
        for r in reversed(refs):
            sp.insert(0, copy.deepcopy(r))

    @staticmethod
    def _sect_is_landscape(sp) -> bool:
        """sectPr 的页面方向判定：orient 属性或宽>高（兼容手工横排未写 orient）。"""
        pgSz = sp.find(qn("w:pgSz"))
        if pgSz is None:
            return False
        if pgSz.get(qn("w:orient")) == "landscape":
            return True
        w, h = pgSz.get(qn("w:w")), pgSz.get(qn("w:h"))
        try:
            return bool(w and h and int(w) > int(h))
        except (TypeError, ValueError):
            return False

    def _find_region(self, title: str):
        """返回 (标题段元素, 收尾sectPr段元素|None)。None＝末节（致谢，body 级 sectPr）。

        区域＝本单元标题段起，至**下一单元标题段**前（用户可能在单元内插入
        横排等自定义节，不能停在第一个 sectPr 段）；收尾段＝区间内最后一个
        带 sectPr 的段（承载本单元节属性，须保留，横排等中段节随内容重建）。
        """
        title_els = [p._p for p in self.doc.paragraphs
                     if p.style and p.style.name in self._TITLE_STYLES and p.text.strip()]
        anchor = None
        want = self._norm_title(title)
        for el in title_els:
            if self._norm_title(self._para_text(el)) == want:
                anchor = el
                break
        if anchor is None:
            # 前缀回退：用户手工把长章题拆成两段时（如第5章），首段命中即可
            for el in title_els:
                t = self._norm_title(self._para_text(el))
                if len(t) > 4 and want.startswith(t):
                    anchor = el
                    break
        if anchor is None:
            raise ValueError(
                f"docx 中找不到单元标题：{title!r}；现存标题：{self.available_titles()}")
        known = getattr(self, "_known_titles", None)
        close = None
        for el in anchor.itersiblings():
            if any(el is t for t in title_els):
                # 下一单元标题为界；拆行的长章题首段（前缀命中已知标题）也是边界
                tnorm = self._norm_title(self._para_text(el))
                if known is None or tnorm in known or any(
                        len(tnorm) > 4 and kt.startswith(tnorm) for kt in known):
                    break
                continue
            if el.tag == qn("w:p"):
                pPr = el.find(qn("w:pPr"))
                if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
                    close = el      # 区间内最后一个 sectPr 段＝本单元收尾
            elif el.tag == qn("w:sectPr"):
                break               # body 级 sectPr（末节致谢）
        return anchor, close

    def available_titles(self) -> list[str]:
        return [p.text.strip() for p in self.doc.paragraphs
                if p.style and p.style.name in self._TITLE_STYLES and p.text.strip()]

    # ── 文本与表格工具 ───────────────────────────────────────
    @staticmethod
    def _para_text(el) -> str:
        return "".join(t.text or "" for t in el.iter(qn("w:t")))

    def _caption_matches(self, el, prefix: str) -> bool:
        if el.tag != qn("w:p"):
            return False
        norm_ = self._para_text(el).replace("　", " ").replace(" ", "")
        return norm_.startswith(prefix)

    def _next_table(self, el):
        """caption 段之后最近的 tbl（空段跳过；遇任何非空段落即止）。"""
        for nxt in el.itersiblings():
            if nxt.tag == qn("w:tbl"):
                return nxt
            if nxt.tag == qn("w:p") and self._para_text(nxt).strip():
                return None
        return None

    def _is_frozen(self, prefix: str | None) -> bool:
        if not prefix:
            return False
        plain = prefix.replace("附表", "表")
        return any(prefix == f or plain == f.replace("附表", "表")
                   for f in self.frozen_tables)

    @staticmethod
    def _tbl_join(tbl_el) -> str:
        """docx 表全部非空有效格文字拼接（前缀缺失时的内容键／冻结表对比用）。"""
        vals = []
        for tc in tbl_el.iter(qn("w:tc")):
            tcPr = tc.find(qn("w:tcPr"))
            vm = tcPr.find(qn("w:vMerge")) if tcPr is not None else None
            if vm is not None and vm.get(qn("w:val")) in (None, "continue"):
                continue
            t = _tc_text(tc)
            if _norm_cell(t):
                vals.append(t)
        return "|".join(vals)

    def _tbl_differs(self, md_el: dict, tbl_el) -> bool:
        """冻结表内容是否与 md 有差异（只读对比，用于告警）。"""
        md_cells = [c for c in (list(md_el["headers"])
                                + [c for r in md_el["rows"] for c in r])
                    if _norm_cell(c)]
        dx = self._tbl_join(tbl_el)
        dx_cells = [c for c in dx.split("|") if _norm_cell(c)]
        if len(md_cells) != len(dx_cells):
            return True
        return any(_norm_cell(a) != _norm_cell(b) for a, b in zip(md_cells, dx_cells))

    # ── 图片哈希（内容级比对键）───────────────────────────────
    @staticmethod
    def _img_hash_md(src: str) -> str | None:
        try:
            return hashlib.sha1(Path(src).read_bytes()).hexdigest()
        except OSError:
            return None

    def _img_hash_dx(self, p_el) -> str | None:
        """docx 图片段 → 图片 part 字节哈希（python-docx 嵌入即原文件字节）。"""
        rid = None
        for blip in p_el.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
            if rid:
                break
        if not rid:
            for im in p_el.iter(_VML_IMAGEDATA):
                rid = im.get(qn("r:id"))
                if rid:
                    break
        if not rid:
            return None
        try:
            part = self.doc.part.rels[rid].target_part
        except KeyError:
            return None
        blob = getattr(part, "blob", None)
        return hashlib.sha1(blob).hexdigest() if blob else None

    # ── 图片版式继承（尺寸＋对齐以 docx 手工精调版为准）───────────
    @staticmethod
    def _img_para_layout(p_el) -> dict | None:
        """图片段落 → 版式（pPr 深拷贝＋首个 drawing 的 cx/cy）；非图片段返回 None。"""
        ext = next(iter(p_el.iter(qn("wp:extent"))), None)
        if ext is None:
            return None
        ppr = p_el.find(qn("w:pPr"))
        return {"pPr": copy.deepcopy(ppr) if ppr is not None else None,
                "size": (ext.get("cx"), ext.get("cy"))}

    def _inherit_img_layout(self, new_p, layout: dict):
        """新渲染图片段继承旧版式：pPr 整体替换（样式/间距/对齐/缩进一次到位），
        wp:extent 与 pic 层 a:ext 同步写回旧尺寸（只改一处 Word 会拉伸）。"""
        if layout.get("pPr") is not None:
            old_ppr = new_p.find(qn("w:pPr"))
            if old_ppr is not None:
                new_p.replace(old_ppr, copy.deepcopy(layout["pPr"]))
            else:
                new_p.insert(0, copy.deepcopy(layout["pPr"]))
        cx, cy = layout["size"]
        if not cx or not cy:
            return
        for e in new_p.iter(qn("wp:extent")):
            e.set("cx", cx)
            e.set("cy", cy)
        for e in new_p.iter(qn("a:ext")):
            e.set("cx", cx)
            e.set("cy", cy)

    def _extract_img_layouts(self, region: list) -> list[dict]:
        """整单元重建前：按文档序收集区域内存图片段落的版式。"""
        return [lay for el in region
                if el.tag == qn("w:p")
                for lay in [self._img_para_layout(el)] if lay is not None]

    def _inherit_img_layouts(self, new_elements: list, saved: list) -> tuple[int, int]:
        """整单元重建后：新图片段按序继承旧版式（图片顺序 md 与 docx 一致）。
        返回 (继承张数, 新图片段总数)；数量差由调用方告警。"""
        n = 0
        imgs = [el for el in new_elements
                if el.tag == qn("w:p") and self._img_para_layout(el) is not None]
        for p, lay in zip(imgs, saved):
            self._inherit_img_layout(p, lay)
            n += 1
        return n, len(imgs)

    # ── 段落级差分：项序列 ───────────────────────────────────
    def _md_items(self, unit: dict) -> list[dict]:
        """合并稿单元 → 比对项 [{kind,key,el,prefix?}]；chapter 元素由标题逻辑处理。"""
        items: list[dict] = []
        last_cap = None
        for el in unit["elements"]:
            t = el["type"]
            if t == "chapter":
                continue
            if t in ("h1", "h2", "h3", "body", "source_note", "list_item",
                     "blockquote"):
                items.append({"kind": "p", "key": norm(el["text"]), "el": el})
            elif t in ("caption_fig", "caption_tab"):
                items.append({"kind": "p",
                              "key": norm(el["text"] + el.get("rest", "")), "el": el})
                if t == "caption_tab":
                    m = _TABLE_PREFIX_RE.match(
                        el["text"].replace(" ", "").replace("　", ""))
                    last_cap = m.group(1) if m else last_cap
            elif t == "keywords_cn":
                items.append({"kind": "p",
                              "key": norm("关键词：" + "；".join(el["keywords"])),
                              "el": el})
            elif t == "keywords_en":
                items.append({"kind": "p",
                              "key": norm("Key Words: " + ", ".join(el["keywords"])),
                              "el": el})
            elif t == "image":
                h = self._img_hash_md(el["src"])
                items.append({"kind": "img",
                              "key": f"img:{h}" if h else f"img-missing:{el['src']}",
                              "el": el})
            elif t == "table":
                key = (f"tbl:{last_cap}" if last_cap else
                       "tbl#" + norm("|".join(
                           c for c in (list(el["headers"])
                                       + [c for r in el["rows"] for c in r])
                           if _norm_cell(c))))
                items.append({"kind": "tbl", "key": key, "el": el,
                              "prefix": last_cap})
                last_cap = None
            elif t == "display_math":
                # 公式段：docx 侧为 OMML（无可比文本），按「公式存在＋编号文本」配对；
                # latex 内容改动不触发差分（改公式请用 --full-unit 或整章重建）
                items.append({"kind": "p",
                              "key": "math|" + norm(f"({el['eq_num']})"
                                                    if el.get("eq_num") else ""),
                              "el": el})
            # footnote / mermaid_image：引擎本不渲染，不参与
        return items

    def _dx_items(self, region: list) -> list[dict]:
        """docx 区域 → 比对项 [{kind,key,el,prefix?,sect_break?}]。

        不参与比对的元素（一律原样保留）：空节界段（横排开收界）、空段/间距段、
        标题残段（手工拆行的长标题后半，由标题逻辑处理）、bookmarkStart 等非段表元素。
        含文字的节界段（Word 存盘把 sectPr 并入段末内容段）作为内容项参与比对并打
        sect_break 标记——保留即原样不动；被变更触及时先把 sectPr 剥离成空段归一。
        """
        items: list[dict] = []
        cur_tag = None
        for el in region[1:]:                    # region[0] 为单元标题段
            if el.tag == qn("w:tbl"):
                key = (f"tbl:{cur_tag}" if cur_tag else
                       "tbl#" + norm(self._tbl_join(el)))
                items.append({"kind": "tbl", "key": key, "el": el, "prefix": cur_tag})
                continue
            if el.tag != qn("w:p"):
                continue
            has_sect = self._sectPr_of(el) is not None
            p = Paragraph(el, self.doc)
            text = p.text
            if el.findall(".//" + qn("m:oMath")):
                # 公式段（OMML）：与 md 侧 display_math 同键配对（"公式存在＋编号文本"）
                it = {"kind": "p", "key": "math|" + norm(text), "el": el}
                if has_sect:
                    it["sect_break"] = True
                items.append(it)
                continue
            m_eq = re.match(r"^\[Equation: .*\](?: \(([^()]*)\))?$", text)
            if m_eq:
                # pandoc 缺失时引擎落回 [Equation: …] 文本，同构 "math|编号" 键配对
                it = {"kind": "p",
                      "key": "math|" + norm(f"({m_eq.group(1)})" if m_eq.group(1) else ""),
                      "el": el}
                if has_sect:
                    it["sect_break"] = True
                items.append(it)
                continue
            if not text.strip():
                if has_sect:
                    continue                     # 空节界段：结构保留
                h = self._img_hash_dx(el)
                if h:
                    items.append({"kind": "img", "key": f"img:{h}", "el": el})
                continue                         # 空段/间距段：保留不动
            if has_sect:
                items.append({"kind": "p", "key": norm(text), "el": el,
                              "sect_break": True})
                continue
            sname = p.style.name if p.style else ""
            if sname in self._TITLE_STYLES:
                continue                         # 标题残段：标题逻辑处理
            if sname.lower().startswith("toc"):
                continue                         # 目录域条目：永不触碰（防御）
            if sname == STYLE["tab_caption"]:
                m = _TABLE_PREFIX_RE.match(text.replace("　", "").replace(" ", ""))
                if m:
                    cur_tag = m.group(1)
            items.append({"kind": "p", "key": norm(text), "el": el})
        return items

    # ── 段落级差分：横排集合（结构变化判据）───────────────────
    def _md_landscape(self, unit: dict) -> set:
        out = set()
        for el in unit["elements"]:
            if el["type"] == "caption_tab" and self.r._is_landscape_caption(el["text"]):
                m = _TABLE_PREFIX_RE.match(el["text"].replace(" ", "").replace("　", ""))
                if m:
                    out.add(m.group(1))
        return out

    def _dx_landscape(self, region: list, anchor_close=None, sentinel=None) -> set:
        """区域内横排表题注前缀集合：表之后最近一个 sectPr 段为横排即横排表
        （横排节界段承载的是它所结束的那个节的属性）。

        单元收尾界段（anchor_close，末节为 body sentinel）同样参与：其属性描述
        单元最后一节——手工横排常以收尾段直接收界（表C.1 实例：C.3+表+C.4
        整体一个横排节，收尾段即横排）。
        """
        seq = list(region[1:])
        tail = anchor_close if anchor_close is not None else sentinel
        if tail is not None and (not seq or seq[-1] is not tail):
            seq.append(tail)
        out, cur_tag, pending = set(), None, None
        for el in seq:
            if el.tag == qn("w:tbl"):
                pending = cur_tag
                continue
            if el.tag != qn("w:p"):
                continue
            sp = self._sectPr_of(el)
            if sp is not None:
                if self._sect_is_landscape(sp) and pending:
                    out.add(pending)
                pending = None
                continue
            p = Paragraph(el, self.doc)
            if p.style and p.style.name == STYLE["tab_caption"]:
                m = _TABLE_PREFIX_RE.match(
                    self._para_text(el).replace("　", "").replace(" ", ""))
                if m:
                    cur_tag = m.group(1)
        return out

    # ── 段落级差分：标题处理 ─────────────────────────────────
    def _title_state(self, anchor_title, title: str) -> tuple[bool, list]:
        """(标题是否实质变化, 手工拆行残段列表)。

        锚点段与连续残段拼接后与 md 标题一致 → 视为未变（保留手工拆行原样）；
        残段防御：标题样式段若命中其他单元已知标题则不认作残段。
        """
        want = self._norm_title(title)
        frags = []
        known = getattr(self, "_known_titles", None) or set()
        for sib in anchor_title.itersiblings():
            if sib.tag != qn("w:p"):
                break
            p = Paragraph(sib, self.doc)
            sname = p.style.name if p.style else ""
            if sname not in self._TITLE_STYLES or not p.text.strip():
                break
            fnorm = self._norm_title(p.text)
            if fnorm in known or any(len(fnorm) > 4 and kt.startswith(fnorm)
                                     for kt in known):
                break                            # 下一单元标题（防御，正常到不了）
            frags.append(sib)
        combined = self._norm_title(self._para_text(anchor_title)) + "".join(
            self._norm_title(self._para_text(f)) for f in frags)
        return (combined != want), frags

    def _rewrite_title(self, anchor_title, title: str, unit_kind: str):
        """标题段原位重写：pPr 与样式保留，只重写 run（<br>→软换行）。"""
        p = Paragraph(anchor_title, self.doc)
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        _fill_title_text(p, _fullwidth_sep(title))
        if unit_kind == "abstract_en":
            from format_engine import FONT_ABSTRACT_TITLE_EN
            for run in p.runs:
                run.font.name = FONT_ABSTRACT_TITLE_EN

    def _unit_header_carrier(self, anchor_title, anchor_close, sentinel):
        """单元页眉载体＝区域内首个带 sectPr 的段（末节为 body sentinel）。"""
        for el in anchor_title.itersiblings():
            if self._sectPr_of(el) is not None:
                return el
        return anchor_close if anchor_close is not None else sentinel

    # ── 段落级差分：应用 ─────────────────────────────────────
    def _render_new(self, el: dict, unit: dict, warnings: list) -> list:
        """单元素渲染（不经横排状态机——增量模式结构不变）。返回新增顶层节点。

        渲染失败时清掉半成品、返回空列表，由调用方决定保留旧元素或跳过插入。
        """
        body = self.doc.element.body
        before = set(body)
        try:
            if unit["kind"] == "references" and el["type"] == "body":
                self.r._add_reference_entry(el["text"])
            else:
                self.r._appendix = (unit["kind"] == "appendix")
                self.r._render_one(el)
                self.r._appendix = False
        except Exception as e:                   # noqa: BLE001 - 失败须保旧内容
            for x in [x for x in body if x not in before]:
                body.remove(x)
            warnings.append(f"元素渲染失败（{el['type']}）：{e}")
            return []
        return [x for x in body if x not in before]

    def _insert_anchor(self, dx_items: list, i2: int, anchor_close, sentinel):
        """插入锚点＝新内容应位于其前的元素；i2 越界则插到单元收尾段/sentinel 前。

        锚点紧贴**纵排**节界段之后时（横排块的开界段，其 sectPr 描述前一节），
        新内容属于前一节，须插到界段之前；横排节界段（收界）之后内容属后一节，不回退。
        """
        if i2 < len(dx_items):
            base = dx_items[i2]["el"]
        else:
            base = anchor_close if anchor_close is not None else sentinel
        if base is anchor_close or base is sentinel:
            return base
        prev = base.getprevious()
        if prev is not None and prev.tag == qn("w:p"):
            sp = self._sectPr_of(prev)
            if sp is not None and not self._sect_is_landscape(sp):
                return prev
        return base

    def _sync_tbl_item(self, m: dict, d: dict, stats: dict, warnings: list,
                       unit: dict, dry_run: bool = False) -> bool:
        """equal 对齐位上的表格：冻结跳过／文字原位同步／结构变化整表重渲染。"""
        prefix = d.get("prefix") or (m.get("prefix") if m else None) or "表?"
        if self._is_frozen(prefix):
            if m and self._tbl_differs(m["el"], d["el"]):
                warnings.append(
                    f"🧊 冻结表 {prefix} 与 md 有差异，按裁定保留 docx 原样")
            stats["kept"] += 1
            return False
        if not m or m["kind"] != "tbl":
            stats["kept"] += 1
            return False
        ok, warn, ncells = _sync_table_text(d["el"], m["el"], prefix,
                                            dry_run=dry_run)
        warnings.extend(warn)
        if ok:
            if ncells:
                stats["tbl_synced"] += 1
                return not dry_run
            stats["kept"] += 1
            return False
        warnings.append(f"{prefix} 结构变化（增删行列），已整表重渲染、格式需手工重做")
        if dry_run:
            stats["replaced"] += 1
            return False
        news = self._render_new(m["el"], unit, warnings)
        if news:
            for e in news:
                d["el"].addprevious(e)
            d["el"].getparent().remove(d["el"])
            stats["replaced"] += 1
            return True
        stats["kept"] += 1
        return False

    def _apply_replace(self, m: dict, d: dict, stats: dict, warnings: list,
                       unit: dict):
        if d["kind"] == "tbl" and self._is_frozen(d.get("prefix")):
            warnings.append(
                f"🧊 冻结表 {d['prefix']} md 侧有变更，按裁定保留 docx 原样"
                "（如需更新请手工处理）")
            stats["kept"] += 1
            return
        img_layout = None
        if d["kind"] == "img" and m["kind"] == "img":
            img_layout = self._img_para_layout(d["el"])  # 移除前收集旧版式
        news = self._render_new(m["el"], unit, warnings)
        if not news:                             # 渲染失败：保留旧元素
            warnings.append(f"保留旧内容：{d['key'][:40]}")
            stats["kept"] += 1
            return
        if img_layout is not None:
            for e in news:
                if e.tag == qn("w:p") and self._img_para_layout(e) is not None:
                    self._inherit_img_layout(e, img_layout)
                    break
        for e in news:
            d["el"].addprevious(e)
        d["el"].getparent().remove(d["el"])
        if d["kind"] == "img" and m["kind"] == "img":
            warnings.append("图片按 md 源文件重渲染；尺寸与对齐已继承 docx 旧版式"
                            "（以手工精调版为准）")
        elif d["kind"] == "tbl" and m["kind"] == "tbl":
            warnings.append(f"{d.get('prefix') or '表格'} 整表重渲染（表号变化或"
                            "题注缺失），格式需手工复核")
        stats["replaced"] += 1

    def _apply_delete(self, d: dict, stats: dict, warnings: list):
        if d["kind"] == "tbl" and self._is_frozen(d.get("prefix")):
            warnings.append(
                f"🧊 冻结表 {d['prefix']} 在 md 中已不存在，按裁定保留 docx 原样")
            stats["kept"] += 1
            return
        d["el"].getparent().remove(d["el"])
        stats["deleted"] += 1

    def _apply_insert(self, m: dict, anchor_el, stats: dict, warnings: list,
                      unit: dict):
        news = self._render_new(m["el"], unit, warnings)
        if not news:
            warnings.append(f"渲染失败，跳过插入：{m['key'][:40]}")
            return
        for e in news:
            anchor_el.addprevious(e)
        if m["kind"] == "tbl":
            warnings.append(f"新增表 {m.get('prefix') or ''} 按基线格式渲染，"
                            "格式需手工精调")
        elif m["kind"] == "img":
            warnings.append("新增图片按默认版式渲染（宽 5.5 英寸＋居中），"
                            "如需调整请在 Word 中精调")
        stats["inserted"] += 1

    def _dispatch_block(self, mspan: list, dspan: list, anchor_el,
                        stats: dict, warnings: list, unit: dict):
        """一个 replace/delete/insert 区间：位置配对，同键按保留处理。"""
        for k in range(max(len(mspan), len(dspan))):
            d = dspan[k] if k < len(dspan) else None
            m = mspan[k] if k < len(mspan) else None
            if d is not None and m is not None:
                if d["key"] == m["key"]:
                    # 同键误配（重复段落对齐漂移）：按未变处理
                    if d["kind"] == "tbl":
                        self._sync_tbl_item(m, d, stats, warnings, unit)
                    else:
                        stats["kept"] += 1
                    continue
                self._apply_replace(m, d, stats, warnings, unit)
            elif d is not None:
                self._apply_delete(d, stats, warnings)
            else:
                self._apply_insert(m, anchor_el, stats, warnings, unit)

    def _replace_unit_incremental(self, unit: dict, title: str,
                                  anchor_title, anchor_close, region: list,
                                  dry_run: bool) -> dict:
        body = self.doc.element.body
        sentinel = body.find(qn("w:sectPr"))
        warnings: list[str] = []
        stats = {"kept": 0, "replaced": 0, "inserted": 0, "deleted": 0,
                 "tbl_synced": 0}

        title_changed, frags = self._title_state(anchor_title, title)
        md_items = self._md_items(unit)
        dx_items = self._dx_items(region)

        def _diff():
            return SequenceMatcher(None,
                                   [d["key"] for d in dx_items],
                                   [m["key"] for m in md_items],
                                   autojunk=False).get_opcodes()

        opcodes = _diff()

        # 变更区间合并（相邻非 equal 操作并为一段），拼接一致区间整体保留
        def _classify(ops):
            keep, work = [], []
            for (di1, di2, mj1, mj2) in self._change_spans(ops):
                dspan, mspan = dx_items[di1:di2], md_items[mj1:mj2]
                if dspan and mspan and \
                        all(i["kind"] == "p" for i in dspan + mspan) and \
                        norm("".join(i["key"] for i in mspan)) == \
                        norm("".join(i["key"] for i in dspan)):
                    keep.append((di1, di2, mj1, mj2))
                else:
                    work.append((di1, di2, mj1, mj2))
            idx = {di for di1, di2, _, _ in keep for di in range(di1, di2)}
            return keep, work, idx

        keep_spans, work_spans, keep_idx = _classify(opcodes)

        # 含文节界段（Word 存盘把 sectPr 并入段末内容段）被变更触及时，先把 sectPr
        # 剥离成其后新空段（结构归一），旧段转为普通内容段；未触及则原样保留。
        # 拼接一致的保留区间不算"触及"——其内节界段一并不动
        split_breaks = 0
        todo = self._breaks_to_split(
            [op for op in opcodes if op[0] == "equal" or op[1] not in keep_idx],
            dx_items, anchor_close, sentinel)
        if todo:
            split_breaks = len(todo)
            if not dry_run:
                for k in todo:
                    new_p = self._split_break_para(dx_items[k]["el"])
                    if dx_items[k]["el"] is anchor_close:
                        anchor_close = new_p
                dx_items = self._dx_items(region)
                opcodes = _diff()
                keep_spans, work_spans, keep_idx = _classify(opcodes)

        if dry_run:
            plan = []
            if title_changed:
                plan.append(f"标题更新：「{self._para_text(anchor_title)[:40]}」→「{title[:40]}」")
            if split_breaks:
                plan.append(f"含文节界段剥离归一 {split_breaks} 段")
            for di1, di2, _, _ in keep_spans:
                plan.append(f"保留手工拆行（拼接一致，{di2 - di1} 段）："
                            f"「{dx_items[di1]['key'][:36]}…」")
            for tag, i1, i2, j1, j2 in opcodes:
                if tag == "equal":
                    for k in range(i2 - i1):
                        d, m = dx_items[i1 + k], md_items[j1 + k]
                        if d["kind"] == "tbl":
                            self._sync_tbl_item(m, d, stats, warnings, unit,
                                                dry_run=True)
                        else:
                            stats["kept"] += 1
                    continue
                if i1 in keep_idx:
                    continue                         # 拼接一致区间：已列入保留计划
                for k in range(i2 - i1):
                    d = dx_items[i1 + k]
                    plan.append(f"{tag}: docx「{d['key'][:36]}」")
                for k in range(j2 - j1):
                    m = md_items[j1 + k]
                    plan.append(f"{tag}:  md「{m['key'][:36]}」")
            stats["kept"] += len(keep_idx)
            return {"unit": unit["id"], "title": title, "mode": "incr",
                    "title_changed": title_changed, "split_breaks": split_breaks,
                    "plan": plan, "warnings": warnings,
                    "changed": bool(plan) or stats["tbl_synced"] > 0, **stats}

        changed = False
        if title_changed:
            self._rewrite_title(anchor_title, title, unit["kind"])
            for f in frags:
                warnings.append(
                    f"删除手工拆行标题残段：「{self._para_text(f)[:30]}」")
                body.remove(f)
            self._sync_header(
                self._unit_header_carrier(anchor_title, anchor_close, sentinel),
                title)
            changed = True
        if split_breaks:
            changed = True

        for tag, i1, i2, j1, j2 in opcodes:
            if tag == "equal":
                for k in range(i2 - i1):
                    d, m = dx_items[i1 + k], md_items[j1 + k]
                    if d["kind"] == "tbl":
                        changed |= self._sync_tbl_item(m, d, stats, warnings, unit)
                    else:
                        stats["kept"] += 1
                continue
            if i1 in keep_idx:
                continue                         # 拼接一致区间：docx 原样保留
            anchor_el = self._insert_anchor(dx_items, i2, anchor_close, sentinel)
            self._dispatch_block(md_items[j1:j2], dx_items[i1:i2], anchor_el,
                                 stats, warnings, unit)
            if i1 != i2 or j1 != j2:
                changed = True

        for (di1, di2, _, _) in keep_spans:
            stats["kept"] += di2 - di1
            warnings.append(
                f"保留手工拆行段落 {di2 - di1} 段（与 md 拼接一致）："
                f"「{dx_items[di1]['key'][:30]}…」")

        return {"unit": unit["id"], "title": title, "mode": "incr",
                "title_changed": title_changed, "split_breaks": split_breaks,
                "warnings": warnings, "changed": changed, **stats}

    @staticmethod
    def _change_spans(opcodes) -> list:
        """相邻非 equal 操作合并为变更区间 [(di1,di2,mj1,mj2)]。"""
        spans = []
        for tag, i1, i2, j1, j2 in opcodes:
            if tag == "equal":
                continue
            if spans and spans[-1][1] == i1 and spans[-1][3] == j1:
                spans[-1] = (spans[-1][0], i2, spans[-1][2], j2)
            else:
                spans.append((i1, i2, j1, j2))
        return spans

    @staticmethod
    def _breaks_to_split(opcodes, dx_items, anchor_close, sentinel) -> list[int]:
        """需要剥离归一的含文节界段下标：①被替换/删除直接触及的；②插入锚点的
        docx 前驱是其一时（新内容须落在该节内，不能隔到节界段之后）。"""
        idx_of = {it["el"]: k for k, it in enumerate(dx_items)}
        todo: set[int] = set()
        n = len(dx_items)
        for tag, i1, i2, j1, j2 in opcodes:
            if tag == "equal":
                continue
            for k in range(i1, i2):
                if dx_items[k].get("sect_break"):
                    todo.add(k)
            base = (dx_items[i2]["el"] if i2 < n else
                    (anchor_close if anchor_close is not None else sentinel))
            prev = base.getprevious() if base is not None else None
            k = idx_of.get(prev)
            if k is not None and dx_items[k].get("sect_break"):
                todo.add(k)
        return sorted(todo)

    # ── 整单元重建（--full-unit／结构变化回退）────────────────
    def _extract_frozen(self, region: list) -> dict:
        """区域内全部表格按题注前缀留存深拷贝（不只冻结表——非冻结表也要
        用旧表骨架做文本原位同步，保住用户手工格式）。"""
        saved = {}
        for el in region:
            if el.tag != qn("w:p"):
                continue
            text = self._para_text(el).replace("　", "").replace(" ", "")
            m = _TABLE_PREFIX_RE.match(text)
            if not m:
                continue
            tbl = self._next_table(el)
            if tbl is not None:
                saved[m.group(1)] = copy.deepcopy(tbl)
        return saved

    def _restore_frozen(self, new_elements: list, saved: dict,
                        unit: dict | None = None) -> tuple[int, list[str]]:
        """在新元素序列中按题注锚点处理旧表：
        - 冻结表：原样回置（文字也不动）；
        - 非冻结表：md 侧文字原位同步进旧表骨架后回置（格式保留、只改文字）；
          结构对不上（增删行列）则保留新渲染表并告警。
        返回 (回置张数, 告警列表)。
        """
        md_tables = _md_tables(unit["elements"]) if unit else {}
        warnings: list[str] = []
        n = 0
        for prefix, old_tbl in saved.items():
            for el in new_elements:
                if not self._caption_matches(el, prefix):
                    continue
                new_tbl = self._next_table(el)
                if new_tbl is None:
                    break
                if prefix in self.frozen_tables:
                    pass                                # 冻结：纯保留
                elif prefix in md_tables:
                    ok, warn, _ncells = _sync_table_text(
                        old_tbl, md_tables[prefix], prefix)
                    warnings.extend(warn)
                    if not ok:
                        warnings.append(
                            f"{prefix} 结构变化（增删行列），已用新渲染表、格式需手工重做")
                        break
                else:
                    warnings.append(f"{prefix} 在 md 中未找到对应表，保留旧表未动")
                new_tbl.addprevious(old_tbl)
                new_tbl.getparent().remove(new_tbl)
                n += 1
                break
        return n, warnings

    @staticmethod
    def _split_break_para(para):
        """含文节界段归一：把 sectPr 剥离到其后新空段（节属性原位保留），
        旧段转为普通内容段；返回新节界段。"""
        pPr = para.find(qn("w:pPr"))
        sectPr = pPr.find(qn("w:sectPr"))
        pPr.remove(sectPr)
        new_p = OxmlElement("w:p")
        new_pPr = OxmlElement("w:pPr")
        new_pPr.append(sectPr)
        new_p.append(new_pPr)
        para.addnext(new_p)
        return new_p

    # ── 替换 ─────────────────────────────────────────────────
    def replace_unit(self, unit: dict, dry_run: bool = False,
                     full_unit: bool = False) -> dict:
        title = _unit_title(unit)
        anchor_title, anchor_close = self._find_region(title)
        body = self.doc.element.body

        # Word 存盘后节尾 sectPr 常并入本单元末个内容段（不再有空段）
        close_has_content = False
        if anchor_close is not None:
            close_has_content = bool("".join(
                t.text or "" for t in anchor_close.iter(qn("w:t"))).strip())

        # 区间＝标题段起，至收尾 sectPr 段前（末节则至 body 级 sectPr 前）
        region = []
        for el in anchor_title.itersiblings():
            if el is anchor_close or el.tag == qn("w:sectPr"):
                break
            region.append(el)
        region.insert(0, anchor_title)
        if close_has_content:
            region.append(anchor_close)   # 并入内容的收尾段按普通内容段参与

        if not full_unit:
            sentinel = body.find(qn("w:sectPr"))
            md_ls = self._md_landscape(unit)
            dx_ls = self._dx_landscape(region, anchor_close, sentinel)
            if md_ls == dx_ls:
                return self._replace_unit_incremental(
                    unit, title, anchor_title, anchor_close, region, dry_run)
            reason = (f"横排表集合变化（docx {sorted(dx_ls) or '∅'} ≠ "
                      f"md {sorted(md_ls) or '∅'}）")
        else:
            reason = "--full-unit"

        if dry_run:
            return {"unit": unit["id"], "title": title, "mode": "full",
                    "reason": reason, "removed": len(region), "changed": True}
        info = self._replace_unit_full(unit, title, anchor_title, anchor_close,
                                       region, close_has_content)
        info["reason"] = reason
        return info

    def _replace_unit_full(self, unit: dict, title: str, anchor_title,
                           anchor_close, region: list,
                           close_has_content: bool) -> dict:
        """整单元重建：删区间全部元素→渲染器重渲染（横排节/页眉引用随重建复位）。"""
        body = self.doc.element.body
        saved_frozen = self._extract_frozen(region)
        saved_imgs = self._extract_img_layouts(region)  # 图片版式（尺寸＋对齐）留存

        # 页眉/页脚引用暂存：本单元引用挂在区域内首个带 sectPr 的段（无内界段则收尾段，
        # 末节在 body 级 sentinel）。区域删除会带走载体；且文末渲染时 python-docx
        # add_section() 克隆的是文末 sentinel、并会剥离其引用（横排节因此会带入错误引用）
        # ——故删除前暂存、渲染搬移后统一重挂（2026-10-07 修复）
        pre_sect = [el for el in region
                    if el is not anchor_title and self._sectPr_of(el) is not None]
        saved_refs = self._clone_refs(pre_sect[0] if pre_sect else anchor_close)
        sentinel = body.find(qn("w:sectPr"))
        saved_sentinel_refs = self._clone_refs(sentinel)

        if close_has_content:
            anchor_close = self._split_break_para(anchor_close)

        # 收尾段节属性归一纵排：用户的横排节可能正好落在单元收尾段上
        # （横排由引擎按 .docx-layout.json 配置重建，不依赖保留）
        if anchor_close is not None:
            sp = anchor_close.find(qn("w:pPr")).find(qn("w:sectPr"))
            pgSz = sp.find(qn("w:pgSz")) if sp is not None else None
            if pgSz is not None and pgSz.get(qn("w:orient")) == "landscape":
                pgSz.set(qn("w:w"), "11906")
                pgSz.set(qn("w:h"), "16838")
                if pgSz.get(qn("w:orient")) is not None:
                    del pgSz.attrib[qn("w:orient")]

        for el in region:
            body.remove(el)

        # 在文末渲染单元新内容（python-docx 自动插到 body sectPr 前），
        # 记录新增元素，整体搬到锚点（收尾 sectPr 段）之前；末节无需搬动
        before = set(body)
        if unit["kind"] in ("chapter", "appendix"):
            self.r._add_unit_title(title)
        elif unit["kind"] == "abstract_cn":
            self.r._add_unit_title("摘　要", front=True)
        elif unit["kind"] == "abstract_en":
            para = self.r._add_unit_title("Abstract", front=True)
            for run in para.runs:
                from format_engine import FONT_ABSTRACT_TITLE_EN
                run.font.name = FONT_ABSTRACT_TITLE_EN
        elif unit["kind"] == "references":
            self.r._add_unit_title("参考文献")
        elif unit["kind"] == "acknowledgement":
            self.r._add_unit_title("致　谢")
        elements = unit["elements"]
        if unit["kind"] in ("chapter", "appendix") and elements and \
                elements[0]["type"] == "chapter":
            elements = elements[1:]
        if unit["kind"] == "references":
            for el in elements:
                if el["type"] == "body":
                    self.r._add_reference_entry(el["text"])
        else:
            self.r._appendix = (unit["kind"] == "appendix")  # 附录小节标题去大纲
            self.r._render_elements(elements)
            self.r._appendix = False
        new_elements = [el for el in body if el not in before]
        if anchor_close is not None:
            for el in new_elements:
                body.remove(el)
                anchor_close.addprevious(el)

        restored, tbl_warnings = self._restore_frozen(new_elements, saved_frozen, unit)
        img_n, img_total = self._inherit_img_layouts(new_elements, saved_imgs)
        if saved_imgs:
            tbl_warnings.append(
                f"图片版式继承 {img_n}/{len(saved_imgs)} 张"
                "（尺寸与对齐以 docx 手工精调版为准）")
        if img_total > len(saved_imgs):
            tbl_warnings.append(
                f"新增 {img_total - len(saved_imgs)} 张图片按默认版式渲染"
                "（宽 5.5 英寸＋居中），如需调整请在 Word 中精调")

        # 引用重挂：剥掉渲染带来的外来引用（文末 sentinel 克隆），
        # 把暂存的单元引用挂回区域内首个节界段（无则收尾段；末节为 sentinel）
        carrier = self._reattach_refs(title, anchor_close, sentinel,
                                      saved_refs, saved_sentinel_refs)
        # 页眉文字随单元标题更新（引用此时已就位）
        self._sync_header(carrier, title)
        return {"unit": unit["id"], "title": title, "mode": "full",
                "removed": len(region), "inserted": len(new_elements),
                "frozen_restored": restored, "warnings": tbl_warnings,
                "changed": True}

    def _reattach_refs(self, title, anchor_close, sentinel,
                       saved_refs, saved_sentinel_refs):
        """渲染搬移后统一重挂引用，返回单元页眉载体（段；末节返回 sentinel）。

        单元引用（页眉/页脚）应挂在区域内**首个**节界段上——该段的节属性描述
        本单元首个节，后续节界段均继承之；挂在靠后的界段会使单元前半段
        继承上一单元的页眉。无内界段时载体＝收尾段，末节＝body sentinel。

        注意：旧 anchor_title 元素在区域删除时已随内容移除，对已脱离文档的元素
        调 itersiblings() 返回空——必须先按标题样式重新定位渲染出的新标题段。
        """
        anchor_title = self._fresh_title(title)
        post = []
        if anchor_title is not None:
            for el in anchor_title.itersiblings():
                if el is anchor_close or el.tag == qn("w:sectPr"):
                    break
                if self._sectPr_of(el) is not None:
                    post.append(el)
        if anchor_close is not None:
            post.append(anchor_close)
        carrier = post[0] if post else anchor_close

        targets = list(post) + ([sentinel] if sentinel is not None else [])
        for el in targets:
            self._strip_refs(el)
        if carrier is None:
            carrier = sentinel
            if not saved_refs:
                saved_refs = saved_sentinel_refs
        if carrier is not None:
            self._attach_refs(carrier, saved_refs)
        if sentinel is not None and sentinel is not carrier:
            self._attach_refs(sentinel, saved_sentinel_refs)
        return carrier

    def _fresh_title(self, title: str):
        """在（已重渲染的）docx 中按标题样式重新定位单元标题段。"""
        want = self._norm_title(title)
        for p in self.doc.paragraphs:
            if p.style and p.style.name in self._TITLE_STYLES and p.text.strip():
                if self._norm_title(p.text) == want:
                    return p._p
        return None

    def _sync_header(self, carrier, title: str):
        """载体（段或 sentinel）所在节的页眉文字更新为单元标题。"""
        target = self._sectPr_of(carrier)
        if target is None:
            return
        target_sec = next((s for s in self.doc.sections if s._sectPr is target), None)
        if target_sec is None or target_sec.header.is_linked_to_previous:
            return
        text = _fullwidth_sep(title).replace("<br>", "")
        for p in target_sec.header.paragraphs:
            if p.text.strip() and p.runs:
                first, rest = p.runs[0], p.runs[1:]
                first.text = text
                for run in rest:
                    run.text = ""
                return

    def save(self, output_path: str | None = None):
        self.doc.save(output_path or self.path)


def backup(docx_path: str) -> str:
    src = Path(docx_path)
    bak_dir = src.parent / "bak"
    bak_dir.mkdir(exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    dst = bak_dir / f"{src.stem}.bak-{ts}{src.suffix}"
    shutil.copy2(src, dst)
    return str(dst)


def _md_tables(elements: list[dict]) -> dict:
    """单元元素中 表题前缀 → 表数据 的映射（md 侧）。"""
    out = {}
    last_cap = None
    for el in elements:
        if el["type"] == "caption_tab":
            m = _TABLE_PREFIX_RE.match(el["text"].replace(" ", "").replace("　", ""))
            last_cap = m.group(1) if m else None
        elif el["type"] == "table":
            if last_cap:
                out[last_cap] = el
            last_cap = None
    return out


def _norm_cell(s: str) -> str:
    """单元格内容级归一（与 norm 同口径，仅少去行首列表符）。"""
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("**", "")
    s = re.sub(r"<br\s*/?>", " ", s, flags=re.IGNORECASE)
    s = re.sub(r"\s*/\s*", "/", s)
    s = re.sub(r"\s+", " ", s)
    s = re.sub(r"\s+(?=[^\x00-\x7f])", "", s)
    s = re.sub(r"(?<=[^\x00-\x7f])\s+", "", s)
    return s.strip()


def _tc_text(tc) -> str:
    return "".join(t.text or "" for t in tc.iter(qn("w:t")))


def _set_tc_text(tc, text: str):
    """把 md 原文写进单元格：保留首段 pPr 与首 run 的 rPr（格式骨架不动）；
    <br> → w:br 断行；用户的多段落单元格在行数一致时按段分配。"""
    text = text.replace("**", "")
    lines = re.split(r"<br\s*/?>", text, flags=re.IGNORECASE)
    ps = tc.findall(qn("w:p"))

    def _write_para(p, line):
        runs = p.findall(qn("w:r"))
        if runs:
            keep = runs[0]
            for child in list(keep):
                if child.tag != qn("w:rPr"):      # 保留 run 级格式
                    keep.remove(child)
            for r in runs[1:]:
                p.remove(r)
        else:
            keep = OxmlElement("w:r")
            p.append(keep)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = line
        keep.append(t)

    if len(ps) == len(lines):                     # 段落数与行数一致：逐段写
        for p, line in zip(ps, lines):
            _write_para(p, line)
    else:                                         # 否则首段内断行
        _write_para(ps[0], lines[0])
        for line in lines[1:]:
            rbr = OxmlElement("w:r")
            rbr.append(OxmlElement("w:br"))
            ps[0].append(rbr)
            rt = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = line
            rt.append(t)
            ps[0].append(rt)
        for p in ps[1:]:
            for r in p.findall(qn("w:r")):
                p.remove(r)


def _sync_table_text(tbl_el, md_el: dict, prefix: str, dry_run: bool = False):
    """md 表格数据原位同步进 docx 表元素（保格式）。返回 (ok, warnings, changed)。

    行/格按"有效单元格"对齐：docx 侧排除纵向合并续格、md 侧排除空占位格。
    结构不一致（增删行列）返回 False，由调用方决定保留新渲染表。
    dry_run=True 只统计不写。
    """
    md_rows = [md_el["headers"]] + md_el["rows"]
    trs = tbl_el.findall(qn("w:tr"))
    warnings: list[str] = []
    if len(trs) != len(md_rows):
        return False, [f"{prefix} 行数不一致（docx {len(trs)} 行 vs md {len(md_rows)} 行）"], 0
    changed = 0
    for tr, mrow in zip(trs, md_rows):
        tcs = []
        for tc in tr.findall(qn("w:tc")):
            tcPr = tc.find(qn("w:tcPr"))
            vm = tcPr.find(qn("w:vMerge")) if tcPr is not None else None
            if vm is not None and vm.get(qn("w:val")) in (None, "continue"):
                continue                          # 纵合并续格：文字在锚点格
            if not _norm_cell(_tc_text(tc)):
                continue                          # 空占位格（分组行/跨年块）不参与
            tcs.append(tc)
        m_vals = [c for c in mrow if _norm_cell(c)]
        if len(tcs) != len(m_vals):
            return False, [f"{prefix} 行结构不一致（docx 有效格 {len(tcs)} vs md {len(m_vals)}）"], 0
        for tc, val in zip(tcs, m_vals):
            if _norm_cell(_tc_text(tc)) != _norm_cell(val):
                if not dry_run:
                    _set_tc_text(tc, val)
                changed += 1
    if changed:
        warnings.append(f"{prefix} 文字原位更新 {changed} 格（格式骨架保留）")
    return True, warnings, changed


def _load_layout(docx_path: str) -> dict:
    """读取 .docx-layout.json（冻结表＋横排表配置）：docx 同目录优先、其 scripts/ 子目录回退。"""
    import json
    for f in (Path(docx_path).parent / ".docx-layout.json",
              Path(docx_path).parent / "scripts" / ".docx-layout.json"):
        if f.exists():
            return json.loads(f.read_text(encoding="utf-8"))
    return {}


def _replacer_layout(docx_path: str) -> dict:
    lay = _load_layout(docx_path)
    return {"frozen_tables": lay.get("frozen", []),
            "landscape_tables": lay.get("landscape", [])}


def main():
    ap = argparse.ArgumentParser(description="docx 章级增量更新（默认段落级最小变更集）")
    ap.add_argument("docx", help="目标论文 docx")
    ap.add_argument("md", help="合并稿（@unit 单元标记）")
    ap.add_argument("--unit", action="append", required=True,
                    help="要替换的单元 id（可多次）：ch1…ch7/abstract_cn/abstract_en/"
                         "references/appendix_a…/acknowledgement")
    ap.add_argument("--full-unit", action="store_true",
                    help="整单元重建（旧行为；段落级差分的兜底路径）")
    ap.add_argument("-o", "--output", default=None, help="输出路径（默认原位更新）")
    ap.add_argument("--no-backup", action="store_true", help="跳过变更前时间戳备份")
    ap.add_argument("--no-purge", action="store_true",
                    help="跳过保存后的未引用媒体清理（默认自动清理）")
    ap.add_argument("--no-stamp", action="store_true",
                    help="跳过保存后的作者元数据改写（默认改为系统用户）")
    ap.add_argument("--dry-run", action="store_true",
                    help="只报告变更计划，不写文件")
    args = ap.parse_args()

    units = {u["id"]: u for u in parse_merged_thesis(args.md)}
    for uid in args.unit:
        if uid not in units:
            print(f"!! 合并稿中无单元 {uid}；可选：{sorted(units)}")
            sys.exit(1)

    replacer = SectionReplacer(args.docx, **_replacer_layout(args.docx))
    replacer._known_titles = {SectionReplacer._norm_title(_unit_title(u))
                              for u in units.values()}
    # 目录单元不在合并稿内（引擎生成的 Word 域），但其标题是真实的单元边界
    replacer._known_titles.add(SectionReplacer._norm_title("目　录"))
    results = []
    for uid in args.unit:
        info = replacer.replace_unit(units[uid], dry_run=args.dry_run,
                                     full_unit=args.full_unit)
        results.append(info)
        if args.dry_run:
            if info["mode"] == "full":
                print(f"[dry-run] {info['unit']}: 整单元重建"
                      f"（{info.get('reason', '')}），将删除 {info['removed']} 个元素（{info['title']}）")
            else:
                marks = f"（{info['title']}）"
                extra = f" · 节界段归一 {info['split_breaks']}" if info.get("split_breaks") else ""
                print(f"[dry-run] {info['unit']}: 保留 {info['kept']} · 替换 {info['replaced']} · "
                      f"插入 {info['inserted']} · 删除 {info['deleted']} · 表同步 {info['tbl_synced']}{extra}{marks}")
                for line in info.get("plan", []):
                    print(f"    {line}")
                for w in info.get("warnings", []):
                    print(f"  ⚠️ {w}")
        else:
            if info["mode"] == "full":
                reason = f"（{info['reason']}）" if info.get("reason") else ""
                extra = f"，表格保留回置 {info['frozen_restored']} 张" if info.get("frozen_restored") else ""
                print(f"OK {info['unit']}: 整单元重建{reason} 删 {info['removed']} 元素 → "
                      f"插 {info['inserted']} 元素（{info['title']}）{extra}")
            else:
                extra = f" · 节界段归一 {info['split_breaks']}" if info.get("split_breaks") else ""
                print(f"OK {info['unit']}: 保留 {info['kept']} · 替换 {info['replaced']} · "
                      f"插入 {info['inserted']} · 删除 {info['deleted']} · 表同步 {info['tbl_synced']}"
                      f"{extra}（{info['title']}）")
            for w in info.get("warnings", []):
                print(f"  ⚠️ {w}")
    if args.dry_run:
        return
    if not any(r.get("changed") for r in results):
        print("全部单元无差异，docx 未写入")
        return
    # 备份紧随写盘前（replacer 全程内存操作，写盘前文件未被触碰）
    if not args.no_backup and args.output is None:
        print(f"备份 → {backup(args.docx)}")
    replacer.save(args.output)
    print(f"已保存：{args.output or args.docx}")
    # 后处理（2026-10-06 用户裁定）：清未引用媒体＋作者改写为系统用户
    from docx_postprocess import purge_unreferenced_media, stamp_author
    target = args.output or args.docx
    if not args.no_purge:
        info = purge_unreferenced_media(target)
        if info["dropped_rels"] or info["dropped_media"]:
            print(f"purge: 删关系 {info['dropped_rels']} 项，删文件 {info['dropped_media']}")
    if not args.no_stamp:
        stamp_author(target)


if __name__ == "__main__":
    main()
