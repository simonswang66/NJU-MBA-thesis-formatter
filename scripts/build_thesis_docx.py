#!/usr/bin/env python3
"""整篇装配 CLI：带单元标记的合并稿 markdown → 学位论文 docx（23E 模板模式）。

用法：
    python3 build_thesis_docx.py <合并稿.md> \
        --template <模板.docx> [--cover <封面模板.docx>] -o <输出.docx>

流程：parse_merged_thesis → Thesis23EEngine 逐单元渲染（Abstract 后插目录域）
      →（可选）merge_cover 前置封面节 → 输出校验摘要。
合并稿单元约定见 SKILL.md「23E 模板模式」节。
"""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from content_parser import parse_merged_thesis
from thesis23e_engine import Thesis23EEngine


def build(md_path: str, template_path: str, output_path: str,
          cover_path: str | None = None) -> dict:
    units = parse_merged_thesis(md_path)
    if not units:
        raise ValueError(f"合并稿中未解析到任何单元：{md_path}")

    # 版式配置（冻结表/横排表）：合并稿同目录的 .docx-layout.json
    import json
    layout_file = Path(md_path).parent / ".docx-layout.json"
    layout = json.loads(layout_file.read_text(encoding="utf-8")) \
        if layout_file.exists() else {}

    engine = Thesis23EEngine(template_path,
                             landscape_tables=layout.get("landscape", []))
    for unit in units:
        engine.render_unit(unit)
        if unit["id"] == "abstract_en":
            engine.add_toc_unit()   # 目录位置：Abstract 之后、正文之前

    # 版式样式归一（图表题/注、文献去自动编号、页眉下边框），作用于样式层
    from docx_postprocess import fix_layout_styles_doc, purge_unreferenced_media, stamp_author
    fix_layout_styles_doc(engine.doc)

    if cover_path:
        tmp = output_path + ".nocover.tmp.docx"
        engine.save(tmp)
        from merge_cover import merge_cover
        merge_cover(tmp, cover_path, output_path)
        os.unlink(tmp)
    else:
        engine.save(output_path)
    # 后处理（2026-10-06 用户裁定）：清未引用媒体 → 作者改写为系统用户
    purge_unreferenced_media(output_path)
    stamp_author(output_path)
    return {"units": [u["id"] for u in units], "output": output_path}


def verify(output_path: str) -> None:
    """重开输出文件，打印结构校验摘要（节、页眉、单元标题、图表计数）。"""
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(output_path)
    print(f"节数：{len(doc.sections)}")
    for i, sec in enumerate(doc.sections):
        hdr = ""
        if not sec.header.is_linked_to_previous:
            hdr = "｜页眉:" + "".join(p.text for p in sec.header.paragraphs if p.text.strip())
        pg = sec._sectPr.find(qn("w:pgNumType"))
        pginfo = ""
        if pg is not None:
            pginfo = f"｜页码:{pg.get(qn('w:fmt'), '')}@{pg.get(qn('w:start'), '续')}"
        print(f"  sec{i}{hdr}{pginfo}")
    titles = [p.text for p in doc.paragraphs
              if p.style and p.style.name in ("MBA论文章标题", "MBA论文前置标题")]
    print("单元标题序列：")
    for t in titles:
        print(f"  {t}")
    print(f"表格 {len(doc.tables)} 张")
    imgs = doc.element.body.findall(f".//{qn('a:blip')}")
    print(f"图片 {len(imgs)} 张")


def main():
    ap = argparse.ArgumentParser(description="合并稿 → 23E 模板 docx 整篇装配")
    ap.add_argument("md", help="带 <!-- @unit:ID --> 单元标记的合并稿路径")
    ap.add_argument("--template", required=True, help="23E 模板 docx 路径")
    ap.add_argument("--cover", default=None, help="封面模板 docx 路径（可选）")
    ap.add_argument("-o", "--output", required=True, help="输出 docx 路径")
    args = ap.parse_args()

    info = build(args.md, args.template, args.output, args.cover)
    print(f"OK 装配 {len(info['units'])} 个单元＋目录 → {info['output']}")
    verify(args.output)


if __name__ == "__main__":
    main()
